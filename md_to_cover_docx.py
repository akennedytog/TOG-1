#!/usr/bin/env python3
"""Convert a cover letter markdown file to a clean Word doc.
Usage: python3 md_to_cover_docx.py <input.md> <output.docx>
"""
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0f, 0x1b, 0x2d)
GOLD = RGBColor(0xc9, 0xa2, 0x27)
GRAY = RGBColor(0x4a, 0x55, 0x68)

def build(in_path, out_path):
    with open(in_path) as f:
        lines = [l.rstrip() for l in f]

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Header
    p = doc.add_paragraph()
    run = p.add_run('ALEC KENNEDY')
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    p = doc.add_paragraph()
    run = p.add_run('Sales Executive Leader')
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = GOLD

    p = doc.add_paragraph()
    run = p.add_run('(502) 403-7201  |  akennedy@theonegroup.info  |  Fort Lauderdale, FL  |  linkedin.com/in/alec-kennedy-1b3a7371')
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY

    doc.add_paragraph()

    # Body: skip the markdown title line (starts with #), keep the rest
    for line in lines:
        if line.startswith('# '):
            continue
        if line.startswith('**Subject:'):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('**', ''))
            run.bold = True
            run.font.size = Pt(11)
            continue
        if line.strip() == '':
            doc.add_paragraph()
        elif line.startswith('> '):
            # blockquote -> italic
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            run.italic = True
        elif line.startswith('---'):
            continue
        else:
            doc.add_paragraph(line)

    doc.save(out_path)
    print("Wrote", out_path)

if __name__ == '__main__':
    build(sys.argv[1], sys.argv[2])
