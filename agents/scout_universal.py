#!/usr/bin/env python3
from __future__ import annotations

import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

INPUT_DIR = Path('/Users/aleckennedy/.openclaw/workspace/data/sales_reports')
OUTPUT_PATH = Path('/Users/aleckennedy/.openclaw/workspace/data/scout_output/briefing_latest.json')

REPORT_PRIORITY = {
    'pod': 1,
    'incentive': 2,
    'accounts_sold': 3,
    'volume': 4,
    'inventory': 5,
    'unknown': 99,
}

IGNORE_SHEETS = {'sheet3'}


@dataclass
class SheetParseResult:
    report_type: str
    sheet_name: str
    columns: list[str]
    records: list[dict[str, Any]]
    summary: dict[str, Any]
    context: dict[str, Any]
    header_row: int | None


def clean_text(value: Any) -> str:
    if value is None:
        return ''
    text = str(value).strip()
    text = re.sub(r'\s+', ' ', text)
    return text


def slugify(value: str) -> str:
    value = clean_text(value).lower()
    value = re.sub(r'[^a-z0-9]+', '_', value).strip('_')
    return value or 'column'


def normalize_number(value: Any) -> float | int | None:
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float):
            if math.isnan(value) or math.isinf(value):
                return None
            rounded = round(value)
            if abs(value - rounded) < 1e-6:
                return int(rounded)
            return round(value, 6)
        return value
    text = clean_text(value).replace(',', '')
    if text in {'#N/A', 'N/A'}:
        return None
    try:
        num = float(text)
    except ValueError:
        return None
    rounded = round(num)
    if abs(num - rounded) < 1e-6:
        return int(rounded)
    return round(num, 6)


def parse_datetime(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.isoformat()
    text = clean_text(value)
    if not text:
        return None
    for fmt in ('%m/%d/%Y %I:%M %p', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).isoformat()
        except ValueError:
            continue
    return text


def extract_account_id(text: str) -> str | None:
    match = re.search(r'\((\d+)\)\s*$', text)
    if match:
        return match.group(1)
    if text.isdigit():
        return text
    return None


def maybe_name_without_id(text: str) -> str:
    return re.sub(r'\s*\(\d+\)\s*$', '', text).strip()


def row_values(ws, row_idx: int, max_cols: int) -> list[Any]:
    return [ws.cell(row_idx, c).value for c in range(1, max_cols + 1)]


def detect_header_row(ws, max_scan_rows: int = 20) -> tuple[int | None, list[str]]:
    max_cols = min(ws.max_column, 30)
    best_row = None
    best_headers: list[str] = []
    best_score = -1
    key_terms = ['row labels', 'customername1', 'date/time', 'promo #', 'customeriddisplay', 'cases']
    for row_idx in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = row_values(ws, row_idx, max_cols)
        headers = [clean_text(v) for v in values]
        non_empty = [h for h in headers if h]
        if len(non_empty) < 2:
            continue
        score = len(non_empty)
        lowered = [h.lower() for h in non_empty]
        score += sum(3 for term in key_terms if term in lowered)
        score += sum(1 for h in lowered if any(ch.isalpha() for ch in h))
        if score > best_score:
            best_score = score
            best_row = row_idx
            best_headers = headers
    return best_row, best_headers


def classify_report(headers: list[str], sheet_name: str) -> str:
    lowered = {clean_text(h).lower() for h in headers if clean_text(h)}
    sheet_lower = sheet_name.lower()
    if {'date/time', 'promo #', 'promotion type', 'brand'} <= lowered:
        return 'incentive'
    if 'row labels' in lowered and any('pod' in h for h in lowered):
        return 'pod'
    if 'customername1' in lowered and 'cases' in lowered:
        return 'volume'
    if 'customername1' in lowered and 'adsname' in lowered:
        return 'accounts_sold'
    if 'customername1' in lowered and 'salespersonname' in lowered:
        return 'inventory'
    if sheet_lower in {'cpgtable', 'cpgtable_chambord'} or sheet_lower.startswith('mar'):
        return 'incentive'
    return 'unknown'


def build_columns(headers: list[str]) -> list[str]:
    seen = Counter()
    columns = []
    for idx, header in enumerate(headers, start=1):
        base = slugify(header) if clean_text(header) else f'column_{idx}'
        seen[base] += 1
        columns.append(base if seen[base] == 1 else f'{base}_{seen[base]}')
    return columns


def iter_records(ws, header_row: int, columns: list[str], max_blank_rows: int = 20) -> Iterable[dict[str, Any]]:
    blank_streak = 0
    for row_idx in range(header_row + 1, ws.max_row + 1):
        values = row_values(ws, row_idx, len(columns))
        if not any(v not in (None, '') for v in values):
            blank_streak += 1
            if blank_streak >= max_blank_rows:
                break
            continue
        blank_streak = 0
        record = {columns[i]: values[i] for i in range(len(columns))}
        yield record


def parse_pod(records: list[dict[str, Any]], columns: list[str]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    out = []
    numeric_cols = [c for c in columns if c != 'row_labels']
    for rec in records:
        owner = clean_text(rec.get('row_labels'))
        if not owner or owner.lower() == 'grand total':
            continue
        metrics = {c: normalize_number(rec.get(c)) for c in numeric_cols}
        if all(v is None for v in metrics.values()):
            continue
        out.append({'owner': owner, 'metrics': {k: v for k, v in metrics.items() if v is not None}})
    totals = defaultdict(float)
    for rec in out:
        for key, value in rec['metrics'].items():
            totals[key] += value
    return out, {'record_count': len(out), 'totals': dict(totals)}


def parse_volume(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    out = []
    total_cases = 0.0
    for rec in records:
        account_name = clean_text(rec.get('customername1'))
        cases = normalize_number(rec.get('cases'))
        if not account_name or cases is None:
            continue
        account_id = extract_account_id(clean_text(rec.get('customeriddisplay')) or account_name)
        out.append({
            'account_name': maybe_name_without_id(account_name),
            'account_id': account_id,
            'cases': cases,
        })
        total_cases += float(cases)
    return out, {'record_count': len(out), 'total_cases': round(total_cases, 6)}


def parse_accounts(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    out = []
    revenue_total = 0.0
    owners = Counter()
    for rec in records:
        account_name = clean_text(rec.get('customername1'))
        ads = clean_text(rec.get('adsname'))
        if not account_name:
            continue
        revenue = normalize_number(rec.get('revenue_net'))
        account_id = extract_account_id(clean_text(rec.get('customeriddisplay')) or account_name)
        out.append({
            'account_name': maybe_name_without_id(account_name),
            'account_id': account_id,
            'ads_name': ads or None,
            'revenue_net': revenue,
        })
        if revenue is not None:
            revenue_total += float(revenue)
        if ads:
            owners[ads] += 1
    return out, {
        'record_count': len(out),
        'revenue_net_total': round(revenue_total, 2),
        'top_ads_by_accounts': owners.most_common(10),
    }


def parse_inventory(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    out = []
    revenue_total = 0.0
    fsm_counts = Counter()
    for rec in records:
        account_name = clean_text(rec.get('customername1'))
        if not account_name:
            continue
        revenue = normalize_number(rec.get('revenue_net'))
        item = {
            'account_name': maybe_name_without_id(account_name),
            'account_id': extract_account_id(clean_text(rec.get('customeriddisplay')) or account_name),
            'salesperson_name': clean_text(rec.get('salespersonname')) or None,
            'field_sales_manager_name': clean_text(rec.get('fieldsalesmanagername')) or None,
            'area_manager_name': clean_text(rec.get('areamanagername')) or None,
            'ads_name': clean_text(rec.get('adsname')) or None,
            'revenue_net': revenue,
        }
        out.append(item)
        if revenue is not None:
            revenue_total += float(revenue)
        if item['field_sales_manager_name']:
            fsm_counts[item['field_sales_manager_name']] += 1
    return out, {
        'record_count': len(out),
        'revenue_net_total': round(revenue_total, 2),
        'top_field_sales_managers': fsm_counts.most_common(10),
    }


def parse_incentive(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    out = []
    promo_types = Counter()
    brands = Counter()
    for rec in records:
        account_name = clean_text(rec.get('dba'))
        promo_type = clean_text(rec.get('promotion_type'))
        brand = clean_text(rec.get('brand'))
        if not account_name:
            continue
        item = {
            'datetime': parse_datetime(rec.get('date_time')),
            'promo_number': clean_text(rec.get('promo')) or clean_text(rec.get('promo_2')) or clean_text(rec.get('promo_')), 
            'photo_taker': clean_text(rec.get('photo_taker')) or None,
            'photo_taker_role': clean_text(rec.get('photo_taker_s_role')) or None,
            'promotion_type': promo_type or None,
            'account_name': account_name,
            'account_id': clean_text(rec.get('account')) or clean_text(rec.get('account_2')) or None,
            'address': clean_text(rec.get('address')) or None,
            'city': clean_text(rec.get('city')) or None,
            'theme': clean_text(rec.get('theme')) or None,
            'elements': clean_text(rec.get('elements')) or None,
            'brand': brand or None,
        }
        out.append(item)
        if promo_type:
            promo_types[promo_type] += 1
        if brand:
            brands[brand] += 1
    return out, {
        'record_count': len(out),
        'promotion_types': promo_types.most_common(),
        'brands': brands.most_common(),
    }


def parse_unknown(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    compact = []
    for rec in records[:100]:
        cleaned = {k: clean_text(v) if isinstance(v, str) else normalize_number(v) if isinstance(v, (int, float)) else v for k, v in rec.items()}
        if any(v not in (None, '', {}) for v in cleaned.values()):
            compact.append(cleaned)
    return compact, {'record_count': len(compact)}


def normalize_headers(headers: list[str]) -> list[str]:
    normalized = []
    for h in headers:
        text = clean_text(h)
        if text == 'Promo #':
            text = 'Promo'
        elif text == "Photo taker's role":
            text = 'Photo Taker Role'
        elif text == 'Account #':
            text = 'Account'
        normalized.append(text)
    return normalized


def parse_sheet(ws) -> SheetParseResult | None:
    if ws.title.lower() in IGNORE_SHEETS:
        return None
    header_row, raw_headers = detect_header_row(ws)
    if not header_row:
        return None
    headers = normalize_headers(raw_headers)
    report_type = classify_report(headers, ws.title)
    columns = build_columns(headers)
    records = list(iter_records(ws, header_row, columns))
    if report_type == 'pod':
        parsed_records, summary = parse_pod(records, columns)
    elif report_type == 'volume':
        parsed_records, summary = parse_volume(records)
    elif report_type == 'accounts_sold':
        parsed_records, summary = parse_accounts(records)
    elif report_type == 'inventory':
        parsed_records, summary = parse_inventory(records)
    elif report_type == 'incentive':
        parsed_records, summary = parse_incentive(records)
    else:
        parsed_records, summary = parse_unknown(records)
    context = extract_context(ws, header_row)
    return SheetParseResult(report_type, ws.title, columns, parsed_records, summary, context, header_row)


def extract_context(ws, header_row: int) -> dict[str, Any]:
    context: dict[str, Any] = {}
    max_cols = min(ws.max_column, 12)
    for row_idx in range(1, header_row):
        vals = row_values(ws, row_idx, max_cols)
        non_empty = [(idx, clean_text(v)) for idx, v in enumerate(vals, start=1) if clean_text(v)]
        if len(non_empty) >= 2:
            key = slugify(non_empty[0][1])
            value = non_empty[1][1]
            if key not in context:
                context[key] = value
        elif len(non_empty) == 1 and 'notes' not in context:
            context['notes'] = non_empty[0][1]
    return context


def parse_workbook(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    report_counts = Counter()
    for ws in wb.worksheets:
        result = parse_sheet(ws)
        if not result or not result.records:
            continue
        report_counts[result.report_type] += 1
        sheets.append({
            'sheet_name': result.sheet_name,
            'report_type': result.report_type,
            'header_row': result.header_row,
            'columns': result.columns,
            'context': result.context,
            'summary': result.summary,
            'records': result.records,
        })
    sheets.sort(key=lambda s: (REPORT_PRIORITY.get(s['report_type'], 99), s['sheet_name'].lower()))
    return {
        'file_name': path.name,
        'sheet_count': len(wb.sheetnames),
        'parsed_sheet_count': len(sheets),
        'report_type_counts': dict(report_counts),
        'sheets': sheets,
    }


def build_briefing(reports: list[dict[str, Any]]) -> dict[str, Any]:
    overall = Counter()
    for report in reports:
        overall.update(report['report_type_counts'])
    return {
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'input_dir': str(INPUT_DIR),
        'report_count': len(reports),
        'overall_report_type_counts': dict(overall),
        'reports': reports,
    }


def main() -> None:
    import sys

    files = sorted(INPUT_DIR.glob('*.xlsx'))
    reports = [parse_workbook(path) for path in files]
    briefing = build_briefing(reports)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(briefing, indent=2), encoding='utf-8')
    sys.stderr.write(f'Wrote {OUTPUT_PATH}\n')
    sys.stderr.write(json.dumps({
        'report_count': briefing['report_count'],
        'overall_report_type_counts': briefing['overall_report_type_counts'],
        'files': [
            {
                'file_name': r['file_name'],
                'parsed_sheet_count': r['parsed_sheet_count'],
                'report_type_counts': r['report_type_counts'],
            }
            for r in reports
        ],
    }, indent=2) + '\n')


if __name__ == '__main__':
    main()
